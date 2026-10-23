"""Build the revised manuscript from its editable Markdown and actual records.
The original 10 October Word draft remains in the campaign backup.
"""
import json,re,shutil
from datetime import datetime,timezone
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parent
ROOT=BASE.parent
RESULTS=ROOT/'gear_inverse_identifiability'/'results'
CAMPAIGN=RESULTS/'corrections_2026-10-12'
OUT=BASE/'MMT_manuscript_corrected.docx'

def rows(name):
    path=RESULTS/(name+'_records.jsonl')
    return [json.loads(l) for l in path.read_text(encoding='utf-8').splitlines() if l.strip()]

def table(doc,headers,data,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in data:
        cells=t.add_row().cells
        for cell,txt in zip(cells,row):cell.text=str(txt)
    for ir,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for ic,cell in enumerate(row.cells):
            cell.width=Cm(widths[ic]);cell.vertical_alignment=1
            tc=cell._tc.get_or_add_tcPr()
            border=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');border.append(el)
            tc.append(border)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');margins.append(el)
            tc.append(margins)
            if ir==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:val'),'clear');shade.set(qn('w:fill'),'E7EDF3');tc.append(shade)
            for p in cell.paragraphs:
                if ir==0:p.paragraph_format.keep_with_next=True
                if ic>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=1.05
                for run in p.runs:run.font.size=Pt(9);run.bold=ir==0
    doc.add_paragraph()
    return t

def paragraph(doc,text,style=None):
    p=doc.add_paragraph(text,style=style)
    p.paragraph_format.widow_control=True
    return p

def figure(doc,name,caption,campaign=False):
    path=(CAMPAIGN if campaign else BASE/'assets')/name
    if not path.exists():
        paragraph(doc,'Figure pending: '+caption)
        return
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(15.5))
    cap=paragraph(doc,caption);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:r.font.size=Pt(9)

def build(source_path=None,out=None,refs_path=None,section_builder='q1_sections',math_revision=6,n_refs=40):
    OUT=Path(out) if out else BASE/'MMT_manuscript_corrected.docx'
    q1=source_path is not None
    doc=Document();sec=doc.sections[0]
    doc.core_properties.author='Romulo Sandoval Flórez; Claudia Patricia Parra Medina; German Leandro Vargas Fonseca'
    doc.core_properties.title='Observable microgeometry and reconstruction ambiguity in spur gears from quasi static transmission error'
    doc.core_properties.subject='Revised working manuscript with audited numerical evidence'
    doc.core_properties.modified=datetime.now(timezone.utc)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.3)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(7)
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.15
    doc.styles['Title'].font.size=Pt(17);doc.styles['Title'].font.bold=True
    for name,size in [('Heading 1',13),('Heading 2',11.5)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.keep_with_next=True
        doc.styles[name].paragraph_format.space_before=Pt(12)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    source=Path(source_path or BASE/'revised_manuscript.md').read_text(encoding='utf-8')
    if q1:
        doc.core_properties.title=source.splitlines()[0].lstrip('# ')
        doc.core_properties.subject='Contact support and practical identifiability research manuscript'
    # Structural citation audit excludes bibliography and cannot substitute
    # for checking whether every cited claim is scientifically supported.
    cited=set()
    for group in re.findall(r'\[([0-9, ]+)\]',source):
        numbers=[int(x) for x in group.replace(' ','').split(',')]
        if 0 in numbers:continue  # [0,1] is a coordinate interval, not a citation.
        cited.update(numbers)
    assert cited==set(range(1,n_refs+1)),('Missing or invalid citations',set(range(1,n_refs+1))-cited)
    for block in source.split('\n\n'):
        text=block.strip()
        if not text:continue
        if text.startswith('# '):paragraph(doc,text[2:],'Title')
        elif text.startswith('### '):doc.add_heading(text[4:],2)
        elif text.startswith('## '):
            doc.add_heading(text[3:],1)
        elif text.startswith('EQ: '):
            if q1:
                from word_math import equation
                number=int(re.search(r'\((\d+)\)\s*$',text).group(1))
                p=paragraph(doc,'');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p._p.append(equation(number,revision=math_revision))
            else:
                p=paragraph(doc,text[4:]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.italic=True;r.font.size=Pt(10)
        elif text.startswith('- '):
            for line in text.splitlines():paragraph(doc,line[2:],'List Bullet')
        elif text.startswith('FIG: '):
            name,caption=text[5:].split(' | ',1);figure(doc,name,caption)
        elif text=='E01_AUDIT':
            path=CAMPAIGN/'E01_verification.json'
            if path.exists():
                rr=json.loads(path.read_text())['records'];base=rr[0]
                paragraph(doc,f"At 120 positions, the maximum force residual is {base['max_force_residual_N']:.3g} N, or {base['relative_force_residual']:.3g} relative to applied force. Normalized complementarity is {base['max_normalized_complementarity']:.3g} and minimum separation is {base['min_separation_mm']:.3g} mm. These support internal numerical consistency within declared tolerances, rather than physical model validation.")
                changes=', '.join(f"{r['relative_rms_change']*100:.4f}% at {r['n_positions']} positions" for r in rr[1:])
                paragraph(doc,'Relative RMS changes under successive doubling are '+changes+'. The last 480-to-960 change meets the 0.1% angular criterion; this does not establish spatial contact convergence or sensitivity convergence.')
        elif text=='E08_SUPPORT_AUDIT':
            path=CAMPAIGN/'E08_contact_support.json'
            if path.exists():
                rr=json.loads(path.read_text())['records']
                lower=min(r['min_loaded_xi'] for r in rr)
                slack=min(r['min_support_separation_mm'] for r in rr)
                change=max(r['max_response_change_mm'] for r in rr)
                paragraph(doc,f"For G2 with nominal addendum relief, the 60-point audits at 50,100 and 300 Nm give minimum loaded xi={lower:.5f}. Candidate contacts with xi<0.30 carry exactly zero computed force and have a separation margin of at least {slack:.4f} mm. Increasing Cd from zero to 0.050 mm at sd=0.30 changes TE by at most {change:.3g} mm. This certifies invisibility for those discrete observations under the specified gap-only model; it does not show that real relief-induced stiffness changes are unobservable.")
        elif text=='TABLE_V6':
            path=CAMPAIGN/'V06_corrected.json'
            if path.exists():
                report=json.loads(path.read_text())
                data=[[r['case'],int(r['torque_Nm']),f"{r['model_pp_um']:.3f}",f"{r['reference_fe_pp_um']:.3f}",f"{r['difference_percent']:+.1f}"] for r in report['results']]
                table(doc,['Geometry','Torque Nm','Model µm','FE µm','Difference %'],data,[4.2,2.7,2.8,2.8,3.9])
                paragraph(doc,'Table 1 Peak to peak comparison with independently checked FE-derived rows of preprint Tables 6-8 [1]. The original 5% target is retained; several cases do not meet it. Surrogate predictions are not validation targets.')
        elif text=='FIG_V6':
            if (CAMPAIGN/'V06_korta_mundo.png').exists():
                figure(doc,'V06_korta_mundo.png','Figure 1 Computed TE waveforms and corrected FE peak to peak windows at 350,500 and 650 Nm. Peak to peak agreement does not establish waveform agreement.',campaign=True)
        elif text=='TABLE_RECORDS':
            data=[]
            for name,expected in [('E04',36),('E07',40),('E08',24)]:
                rec=rows(name)
                data.append([name,len(rec),expected,sum(r.get('success',False) for r in rec),sum('provenance' in r for r in rec)])
            table(doc,['Experiment','Records','Planned','Optimizer success','Versioned new'],data,[3.3,2.4,2.4,4.3,4.0])
            paragraph(doc,'Table 3 Record inventory at manuscript generation. Solver success does not imply physical validation. Historical records without provenance remain unversioned; two pilot cases have a tagged scaled-central derivative policy.')
        elif text=='TABLE_E5':
            path=CAMPAIGN/'E05_fixed_budget.json'
            if not path.exists():paragraph(doc,'Fixed-budget E5 audit pending completion; the historical full-Fisher ranking is withdrawn.')
            else:
                report=json.loads(path.read_text())
                data=[]
                for d in report['designs']:
                    a=d['conditional_addendum'];f=d['full']
                    data.append([','.join(str(int(t)) for t in d['torques']),f['n_observations'],f['rank'],f"{a['logdet']:.4f}",f"{a['min_eigenvalue']:.3g}"])
                table(doc,['Torques Nm','Observations','Full rank','Conditional log det','Conditional E'],data,[3.2,2.9,2.4,4.1,3.8])
                paragraph(doc,'Table 4 Local information at 120 observations, scaled parameters and sigma=0.1 µm. Addendum scores condition on fixed dedendum. Rank tolerance is 1e-7 relative to the largest singular value.')
                audit=report['active_contact_audit']
                paragraph(doc,'Contact audit at 240 positions: '+ '; '.join(f"{r['torque']:.0f} Nm: minimum active xi={r['min_active_profile_coordinate']:.6f}, {r['dedendum_active_samples']} loaded points below sd=0.25" for r in audit)+'. This is sampled evidence, not a continuous-domain structural proof.')
        elif text=='TABLE_E3':
            path=CAMPAIGN/'E03_sampling_audit.json'
            if path.exists():
                report=json.loads(path.read_text())
                table(doc,['Positions','Numerical rank','Root change RMSE µm'],
                      [[r['n_positions'],r['rank'],f"{r['root_change_rmse_um']:.6g}"] for r in report['grids']],
                      [4.0,4.0,8.4])
                paragraph(doc,'Table 2 Sampling audit with nonlinear contact at 100 Nm. Root change sets sd=0.15 while holding other parameters at truth. Relative SVD rank tolerance is 1e-7. Tiny response changes relative to sigma=0.1 µm demonstrate practical weakness rather than zero continuous sensitivity.')
        elif text=='FIG_E3':
            if (CAMPAIGN/'E03_sampling_audit.png').exists():
                figure(doc,'E03_sampling_audit.png','Figure 2 Sampling dependence of numerical rank and fit versus held-out dense-grid response errors.',campaign=True)
        elif text=='FIG_E5':
            if (CAMPAIGN/'E05_fixed_budget.png').exists():figure(doc,'E05_fixed_budget.png','Figure 4 Fixed observation budget and scaled sensitivity spectra. Projected scores are conditional addendum information.',campaign=True)
        elif text=='FIG_PARAMETER_BLOCKS':
            if (CAMPAIGN/'parameter_block_errors.png').exists():
                figure(doc,'parameter_block_errors.png','Figure 3 Addendum and dedendum parameter errors versus response residual in E4/E7. Each point is a stored run; single and multi load have different historical observation counts. This is an exploratory comparison, not a confidence or threshold estimate.',campaign=True)
        elif text.startswith('Q1_'):
            from importlib import import_module
            add_section=import_module(section_builder).add_section
            add_section(doc,text,paragraph,table,figure)
        elif text=='REFERENCES_40':
            refs=json.loads(Path(refs_path or ROOT/'referencias_40_seleccionadas.json').read_text(encoding='utf-8'))
            assert len(refs)==n_refs
            for i,r in enumerate(refs,1):
                authors=r.get('author',[])
                names=', '.join((a.get('given','')+' '+a.get('family','')).strip() for a in authors[:6])
                if len(authors)>6:names+=', et al.'
                title=' '.join(r['title'][0].split())
                venue=(r.get('container-title') or [r.get('publisher','')])[0]
                ref=f"[{i}] {names.rstrip('.')}. {title.rstrip('.')}. {venue} ({r['year']})"
                if r.get('volume'):ref+=', '+r['volume']
                if r.get('page'):ref+=', '+r['page']
                ref+='.'
                if r.get('DOI'):ref+=' doi:'+r['DOI']+'.'
                else:ref+=' '+r['url']
                p=paragraph(doc,ref);p.paragraph_format.space_after=Pt(6)
                for run in p.runs:run.font.size=Pt(10)
        else:paragraph(doc,text)
    # Remove title border residue and theme fonts inherited from the template.
    for root in [doc.element, doc.styles.element]:
        for border in root.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        for fonts in root.xpath('.//w:rFonts'):
            for attribute in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                fonts.attrib.pop(qn('w:'+attribute),None)
    # Put table captions before their table so pagination cannot orphan them.
    for p in list(doc.paragraphs):
        if re.match(r'^Table \d+ ',p.text):
            p.paragraph_format.keep_with_next=True
            if q1:continue  # Q1 captions are already inserted before tables.
            previous=p._p.getprevious()
            while previous is not None and previous.tag!=qn('w:tbl'):
                previous=previous.getprevious()
            if previous is not None:
                previous.addprevious(p._p)
                p.paragraph_format.keep_with_next=True
    if q1:
        from copy import deepcopy
        from docx.text.run import Run
        all_paragraphs=list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
        pattern=re.compile(r'\b([Cs gurtxyβ αξλJ𝓘waF])_([a-z0-9]+)'.replace(' ',''))
        for p in all_paragraphs:
            for run in list(p.runs):
                if not pattern.search(run.text):continue
                segments=[];offset=0
                for match in pattern.finditer(run.text):
                    segments.extend([(run.text[offset:match.start()],False),(match[1],False),(match[2],True)])
                    offset=match.end()
                segments.append((run.text[offset:],False))
                for text,is_sub in segments:
                    if not text:continue
                    element=deepcopy(run._r)
                    newrun=Run(element,p);newrun.text=text
                    if is_sub:newrun.font.subscript=True
                    run._r.addprevious(element)
                run._r.getparent().remove(run._r)
    doc.save(OUT)
    (BASE/'corrected_manuscript_extracted.txt').write_text('\n'.join(p.text for p in doc.paragraphs),encoding='utf-8')
    print('Saved',OUT)
    return OUT

if __name__=='__main__':build()
