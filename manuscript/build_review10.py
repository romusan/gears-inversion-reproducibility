"""Build revision 10 from the v9 markdown: shorten the abstract to 220-250
words with a contribution-ordered structure, cite the two inverse-analysis
references added as [41,42], add the convergence-robustness paragraph to
Section 4.3 (phase 12), scope the finite-element validation claims to the
sampled interior phases, name the AI tool in the declaration, point the data
availability at tag manuscript-v10, and emit the independent highlights file.

Every edit is an asserted string replacement so drift in the source fails
the build instead of patching silently.
"""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from build_corrected_manuscript import build, paragraph, table
from build_review5 import format_inline
import review5_sections as v5
import review10_sections  # noqa: F401  (used by section_builder name)
BASE = Path(__file__).resolve().parent

ABSTRACT = (
    "Transmission-error fitting can conceal uncertainty in gear profile reconstruction "
    "when relief enters only marginally into the loaded flank. A quasi-static model with "
    "four profile parameters establishes a sufficient contact-support condition for exact "
    "relief invisibility; implicit sensitivities and noise experiments distinguish "
    "them from a weak nonzero-sensitivity transition. At the nominal "
    "geometry, median scaled reconstruction errors for a marginal and a deeper dedendum "
    "start are 0.394 and 0.106 over 100 realizations at 120 observations, 300 Nm, 0.1 "
    "micrometre noise; Fisher predictions agree with Monte Carlo in the deeper regime but "
    "fail near invisible sets. The contrast is geometry dependent: comparing the minimum "
    "loaded flank coordinate with the expected relief-start range screens relief before "
    "inversion, complemented by angular-grid refinement and singular values. "
    "Equal-budget controls expose sampling confounding in load "
    "comparisons, and complementary-phase noise experiments show no statistically "
    "resolved two-load advantage. All 48 external comparisons overpredict "
    "peak-to-peak transmission error, and correcting the shear denominator alone worsens "
    "agreement: contact support explains internal observability, but independent "
    "stiffness and waveform evidence remains necessary for physical reconstruction. "
    "Finite-element cycles inverted by the analytical operator confirm exact dedendum "
    "invisibility outside loaded support over the sampled interior phases and recover the "
    "0.020 mm addendum depth to 0.0006 mm; where the dedendum is observable its depth is "
    "recovered within 0.2 micrometre noise-free, its start only coarsely, and the "
    "finite-element relief response is ninefold stronger, corroborating the structural "
    "discrepancy. A non-parametric resolution analysis identifies the relief over the "
    "loaded flank up to an additive constant, with sub-micrometre precision in most "
    "resolved directions."
)

HIGHLIGHTS = [
    "Relief visibility depends on penetration into the loaded part of the flank",
    "Minimum loaded flank coordinate screens relief visibility before inversion",
    "Mean-removed contact observable leaves relief known up to a constant",
    "Non-parametric resolution analysis maps observable and null flank coordinates",
    "External benchmark comparisons remain systematically discrepant",
]


def rep(t, old, new, count=1):
    assert old in t, 'v9 text drifted: ' + old[:80]
    return t.replace(old, new, count)


def main():
    t = (BASE / 'q1_manuscript_v9.md').read_text(encoding='utf-8')

    # --- abstract: 308 -> 220-250 words, contribution-ordered ---------------
    head, rest = t.split('## Abstract\n\n', 1)
    _, tail = rest.split('\n\nKeywords:', 1)
    n_words = len(ABSTRACT.split())
    assert 220 <= n_words <= 250, n_words
    t = head + '## Abstract\n\n' + ABSTRACT + '\n\nKeywords:' + tail

    # --- introduction: inverse-analysis precedents [41,42] --------------------
    t = rep(t,
        "An earlier inverse study estimated meshing TE from tooth-root stresses [8]; its "
        "observation and unknown differ from the microgeometry reconstruction considered "
        "here. This study therefore makes a scoped contribution to numerical "
        "characterization of the inverse operator, rather than claiming that inverse gear "
        "analysis is new or that no related reconstruction method exists.",
        "An earlier inverse study estimated meshing TE from tooth-root stresses [8]. "
        "Analytical inverse solutions of the tooth contact analysis problem also exist "
        "[41], and double-flank roll-test measurements have been inverted to determine "
        "gear parameter deviations [42]; the observation, unknown and model class of "
        "these works differ from the quasi-static microgeometry reconstruction "
        "considered here. This study therefore makes a scoped contribution to numerical "
        "characterization of the inverse operator, rather than claiming that inverse "
        "gear analysis is new or that no related reconstruction method exists.")

    # --- section 4.3: convergence robustness (phase 12) ----------------------
    t = rep(t,
        "Simulation σ values are instrument-noise scenarios, not estimates of the external "
        "model discrepancy.\n\nQ1_PROFILES",
        "Simulation σ values are instrument-noise scenarios, not estimates of the external "
        "model discrepancy.\n\n"
        "Each fit retains its optimizer success status and the statistical summaries "
        "include failed optimizations rather than removing them. To test whether these "
        "flags bias the results, all 65 non-converged fits — 45 of 1200 in the main "
        "campaign and 20 of 600 in the complementary-phase control, spread over nine "
        "cells and concentrated at the highest noise levels — were re-estimated from the "
        "same observations with four additional broad starts and an increased evaluation "
        "budget. Only 14 of the 65 changed their success flag, and the per-cell medians "
        "of the dedendum scaled error move by at most 0.0034 (0.394 to 0.397 in the "
        "nominal cell). The non-convergence flags therefore record optimizer termination "
        "in flat cost regions, not a bias of the identification, and the reported "
        "conclusions are invariant to them.\n\n"
        "Q1_PROFILES")

    # --- discussion and conclusions: scope FE claims to sampled phases --------
    t = rep(t,
        "The cross-operator test of Section 4.5 partially fills that gap: finite-element "
        "responses of known geometry inverted with the analytical operator recover",
        "The cross-operator test of Section 4.5 partially fills that gap: over the "
        "sampled interior phases, finite-element responses of known geometry inverted "
        "with the analytical operator recover")
    t = rep(t,
        "reproduce exact dedendum invisibility outside loaded support, and bound the "
        "residual between the two mechanical operators",
        "reproduce exact dedendum invisibility outside loaded support at those phases, "
        "and bound the residual between the two mechanical operators")
    t = rep(t,
        "no dedendum signal exists to recover outside loaded support, and the "
        "inter-operator residual bounds practical precision",
        "no dedendum signal exists to recover outside loaded support at the sampled "
        "phases, and the inter-operator residual bounds practical precision")

    # --- declarations: name the AI tool ---------------------------------------
    t = rep(t,
        "Declaration of generative AI use. During preparation of this work, the authors "
        "used AI tools to assist with code inspection, numerical analysis and manuscript "
        "drafting. The authors reviewed the content and take full responsibility for the "
        "scientific accuracy and integrity of the article.",
        "Declaration of generative AI use. During preparation of this work, the authors "
        "used Kimi (Moonshot AI) to assist with code inspection, numerical analysis and "
        "manuscript drafting. The authors reviewed and validated all assisted content "
        "and take full responsibility for the scientific accuracy and integrity of the "
        "article.")

    # --- data availability: v9 -> v10 ------------------------------------------
    t = rep(t,
        "https://github.com/romusan/gears-inversion-reproducibility/tree/manuscript-v9. The "
        "annotated tag manuscript-v9 and a SHA-256 manifest identify the deposited version. The "
        "preceding manuscript-v8 tag remains available. No DOI has been assigned.",
        "https://github.com/romusan/gears-inversion-reproducibility/tree/manuscript-v10. The "
        "annotated tag manuscript-v10 and a SHA-256 manifest identify the deposited version. "
        "Earlier manuscript tags remain available. No DOI has been assigned.")

    (BASE / 'q1_manuscript_v10.md').write_text(t, encoding='utf-8')

    v5.SUPPLEMENT.clear()
    path = build(BASE / 'q1_manuscript_v10.md',
                 BASE / 'MMT_manuscript_Q1_revision_v10.docx',
                 BASE.parent / 'referencias_42_q1.json',
                 section_builder='review10_sections', math_revision=7,
                 n_refs=42)
    d = Document(path); format_inline(d)
    for tab_ in d.tables:
        for cell in tab_.rows[-2].cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    d.save(path)

    s = Document(path)
    for child in list(s.element.body):
        if child.tag != qn('w:sectPr'):
            s.element.body.remove(child)
    paragraph(s, 'Supplementary numerical controls for spur gear microgeometry inversion', 'Title')
    paragraph(s, 'Local information, phase-grid controls, finite-element resolution and angular '
                 'sampling for the accompanying manuscript. The original linear FE records and '
                 'the quadratic weak-contact records use different interface discretizations and '
                 'are labelled separately. Tables S10–S13 report the cross-operator inversion '
                 '(nominal and observable-dedendum tests), formulation-robustness and '
                 'resolution-refinement records.')
    for item in sorted(v5.SUPPLEMENT, key=lambda x: int(x['number'][1:])):
        if item['number'] == 'S7':
            item['caption'] = item['caption'].replace(
                'The interface stiffness sequence is nonmonotonic.',
                'The sampled interface stiffness sequence is nonmonotonic across refinement levels.')
        p = paragraph(s, item['caption']); p.paragraph_format.keep_with_next = True
        table(s, item['headers'], item['rows'], item['widths'])
    paragraph(s, 'S1–S4 retain the original statistical and local-information controls. S5–S6 and '
                 'S8 preserve the original linear/nodal FE records. S7 varies quadratic bulk and '
                 'pressure controls separately; the sampled interface stiffness sequence is '
                 'nonmonotonic across refinement levels. S9 reports '
                 'all offsets for the declared interpolants; the eight-node FE waveform is an '
                 'interpolation diagnostic and is not a direct FE angular scan. S10 reports '
                 'cross-operator inversion against the two-dimensional finite-element generator '
                 '(nineteen interior phases; six restricted phases excluded). S11 repeats the '
                 'nominal noise cells under the energy-projected compliance formulation. S12 '
                 'refines the non-parametric resolution analysis over the flank grid and the '
                 'angular sampling. S13 reports the observable-dedendum cross-operator inversion '
                 'on the symmetric 24/24 pair. Monte Carlo campaigns use fixed, disjoint seed '
                 'bases (20261012 in the original campaign, 20261022 in the formulation '
                 'repetition); shared seed indices act as common random numbers within each '
                 'paired comparison.')
    format_inline(s)
    for tab_ in s.tables:
        for cell in tab_.rows[-2].cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    s.save(BASE / 'MMT_supplementary_Q1_revision_v10.docx')

    (BASE / 'highlights.txt').write_text(
        'Highlights\n' + '\n'.join('• ' + h for h in HIGHLIGHTS) + '\n',
        encoding='utf-8')
    print('Saved revision 10, abstract words', n_words,
          ', supplementary tables', len(v5.SUPPLEMENT))


if __name__ == '__main__':
    main()
