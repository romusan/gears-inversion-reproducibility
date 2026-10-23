"""Build the article and supplement from final scientific sources."""
import json,hashlib
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
from build_corrected_manuscript import build,paragraph,table
from build_review5 import format_inline
import review5_sections as v5

BASE=Path(__file__).resolve().parent;ROOT=BASE.parent

def normalize(doc):
    format_inline(doc)
    import re
    from copy import deepcopy
    from docx.text.run import Run
    # Complete the inline subscript vocabulary, including Greek indices and primes.
    pattern=re.compile(r'([A-Za-zα-ωΑ-Ωℓ]′?)_([A-Za-z0-9ℓ]+)')
    paragraphs=list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]
    for p in paragraphs:
        for run in list(p.runs):
            if not pattern.search(run.text):continue
            segments=[];offset=0
            for match in pattern.finditer(run.text):
                segments.extend([(run.text[offset:match.start()],False),(match[1],False),(match[2],True)]);offset=match.end()
            segments.append((run.text[offset:],False))
            for text,is_sub in segments:
                if not text:continue
                element=deepcopy(run._r);new=Run(element,p);new.text=text
                if is_sub:new.font.subscript=True
                run._r.addprevious(element)
            run._r.getparent().remove(run._r)
    # Tables fit the 16.4 cm text area; preserve proportional column widths.
    for tab in doc.tables:
        widths=[c.width for c in tab.columns];total=sum(widths)
        if total>Cm(16.4):
            for col,w in zip(tab.columns,widths):col.width=int(w*Cm(16.4)/total)
            for row in tab.rows:
                for cell,w in zip(row.cells,widths):cell.width=int(w*Cm(16.4)/total)
        for row in tab.rows:
            for cell in row.cells:
                for margin in cell._tc.xpath('./w:tcPr/w:tcMar/w:top | ./w:tcPr/w:tcMar/w:bottom'):margin.set(qn('w:w'),'75')
                for p in cell.paragraphs:
                    for run in p.runs:run.font.size=Pt(8.5)
        # Keep at least two body rows together at each end of a split table.
        for index in [1,len(tab.rows)-2]:
            if index>0:
                for cell in tab.rows[index].cells:
                    for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    # Replace inherited small-label bitmaps with regenerated data-equivalent plots.
    old=ROOT/'gear_inverse_identifiability/results/q1_revision_2026-10-12'
    mapping={ROOT/'gear_inverse_identifiability/results/q1_review2_2026-10-12/forward_ablation.png':'fig_benchmark_v11.png',ROOT/'gear_inverse_identifiability/results/q1_review2_2026-10-12/plateau_stiffness.png':'fig_stiffness_v11.png',old/'support_transition.png':'fig_transition_v11.png',old/'profile_likelihood.png':'fig_profiles_v11.png',old/'monte_carlo.png':'fig_monte_v11.png',ROOT/'revision_5_2026-10-12/analytic_spectra_v5.png':'fig_spectra_v11.png'}
    replacements={hashlib.sha256(k.read_bytes()).hexdigest():(BASE/'assets'/v).read_bytes() for k,v in mapping.items()}
    for rel in doc.part.rels.values():
        if 'image' in rel.reltype:
            digest=hashlib.sha256(rel.target_part.blob).hexdigest()
            if digest in replacements:rel.target_part._blob=replacements[digest]
    from PIL import Image
    import io
    for shape in doc.inline_shapes:
        rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        part=doc.part.related_parts[rid]
        with Image.open(io.BytesIO(part.blob)) as im:w,h=im.size
        width=min(Cm(15.5),int(Cm(18.5)*w/h))
        if hashlib.sha256(part.blob).hexdigest()==hashlib.sha256((BASE/'assets/fig_stiffness_v11.png').read_bytes()).hexdigest():width=Cm(12)
        if hashlib.sha256(part.blob).hexdigest()==hashlib.sha256((BASE/'assets/fig_transition_v11.png').read_bytes()).hexdigest():width=int(Cm(14.5)*w/h)
        shape.width=width;shape.height=int(width*h/w)
    # Float the large profile figure past two short paragraphs in the same
    # uncertainty subsection so a table does not leave half a page empty.
    caption=next((p for p in doc.paragraphs if p.text.startswith('Figure 4 ')),None)
    anchor=next((p for p in doc.paragraphs if p.text.startswith('Box constraints affect')),None)
    if caption is not None and anchor is not None:
        picture=caption._p.getprevious()
        if picture is not None and picture.xpath('.//w:drawing'):
            anchor._p.addnext(picture);picture.addnext(caption._p)

def main():
    v5.SUPPLEMENT.clear()
    common_extract=BASE/'corrected_manuscript_extracted.txt'
    previous_extract=common_extract.read_bytes() if common_extract.exists() else None
    try:
        path=build(BASE/'q1_manuscript_v11.md',BASE/'MMT_manuscript_Q1_revision_v11.docx',ROOT/'referencias_42_q1.json',section_builder='review11_sections',math_revision=7,n_refs=42)
        (BASE/'manuscript_v11_extracted.txt').write_bytes(common_extract.read_bytes())
    finally:
        if previous_extract is not None:common_extract.write_bytes(previous_extract)
    doc=Document(path);normalize(doc);doc.save(path)
    sup=Document(path)
    for child in list(sup.element.body):
        if child.tag!=qn('w:sectPr'):sup.element.body.remove(child)
    paragraph(sup,'Supplementary controls for spur gear microgeometry identifiability','Title')
    paragraph(sup,'Numerical evidence for local information, sampling design, finite-element contact and cross-operator inversion. Tables distinguish the 57/57 continuum controls from the dedicated 24/24 gap test. All precision statements assume the declared observation model; a contact response does not by itself validate geometrical reconstruction.')
    items=sorted(v5.SUPPLEMENT,key=lambda x:int(x['number'][1:]))
    assert [r['number'] for r in items]==['S'+str(i) for i in range(1,17)]
    for item in items:
        p=paragraph(sup,item['caption']);p.paragraph_format.keep_with_next=True
        table(sup,item['headers'],item['rows'],item['widths'])
    paragraph(sup,'S1–S4 concern internal information and sampling. S5–S9 concern the 57/57 FE and angular controls; linear nodal and quadratic weak-contact records are distinct. S10 reports the 57/57 gap inversion and S11 the formulation repetition. S12 describes the local discretized resolution. S13–S16 report the 24/24 archived inversions, broad-start and restricted fits, profile diagnostics and dedicated bulk-mesh study. Their fixed 0.02 mm contact interface and nominal contour do not certify interface refinement or physical material removal. Every selected fit and its optimizer status remain in the numerical data.')
    normalize(sup);sup.save(BASE/'MMT_supplementary_Q1_revision_v11.docx')
    (BASE/'highlights_v11.txt').write_text('Highlights\n\n• Loaded contact support defines exactly invisible relief sets in a fixed-gap model\n• Weak support overlap separates local rank from practical profile recovery\n• Mean-removed transmission error leaves an additive relief gauge unresolved\n• Restricted fits distinguish dedendum response from geometrical reconstruction\n• External amplitude discrepancy limits physical interpretation of inverse fits\n',encoding='utf-8')
    (ROOT/'revision_11_2026-10-23/supplement_tables.json').write_text(json.dumps(items,indent=2,ensure_ascii=False),encoding='utf-8')
    print('Saved manuscript and supplement with',len(items),'supplementary tables')

if __name__=='__main__':main()
