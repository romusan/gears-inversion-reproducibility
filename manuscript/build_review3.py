from pathlib import Path
from prepare_review3 import main as prepare
from build_corrected_manuscript import build
BASE=Path(__file__).resolve().parent
if __name__=='__main__':
    prepare()
    out=build(BASE/'q1_manuscript.md',BASE/'MMT_manuscript_Q1_revision_v3.docx',BASE.parent/'referencias_40_q1.json',section_builder='review3_sections')
    from docx import Document
    doc=Document(out)
    # Keep at least the final two data rows together across page boundaries.
    for t in doc.tables:
        if len(t.rows)>3:
            for cell in t.rows[-2].cells:
                for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    doc.save(out)
