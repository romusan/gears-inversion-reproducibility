from pathlib import Path
from prepare_review3 import main as prepare
from build_corrected_manuscript import build
from docx import Document
BASE=Path(__file__).resolve().parent
if __name__=='__main__':
    prepare();p=BASE/'q1_manuscript.md';t=p.read_text(encoding='utf-8')
    t=t.replace('### 3.2 Implicit contact tangent','Q1_COUPLED_OBSERVATION\n\n### 3.2 Implicit contact tangent')
    t=t.replace('## Declarations','## Appendix A Exploratory continuum and support controls\n\nQ1_FE_APPENDIX\n\n## Declarations')
    t=t.replace('This support diagnostic must be checked against angular refinement and the singular spectrum;',
        'This support diagnostic must be checked across mechanical variants and loads, against angular refinement and the singular spectrum;')
    t=t.replace('external waveform and stiffness data are required to determine whether their shape and plateau levels reproduce the mechanical response.',
        'the exploratory coupled FE contrast supports approximately common compliance between two interior phases but does not reproduce the missing corner-contact mechanics. Additional edge-bearing phases cannot establish full-cycle benchmark agreement.')
    (BASE/'q1_manuscript_v4.md').write_text(t,encoding='utf-8')
    out=build(BASE/'q1_manuscript_v4.md',BASE/'MMT_manuscript_Q1_revision_v4.docx',BASE.parent/'referencias_40_q1.json',section_builder='review4_sections')
    d=Document(out)
    for table in d.tables:
        if len(table.rows)>3:
            for cell in table.rows[-2].cells:
                for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    d.save(out)
