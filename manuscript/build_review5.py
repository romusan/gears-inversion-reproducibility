import re,json
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.run import Run
from docx.oxml.ns import qn
from build_corrected_manuscript import build,paragraph,table
import review5_sections as v5
BASE=Path(__file__).resolve().parent

def format_inline(d):
    paragraphs=list(d.paragraphs)+[p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    pattern=re.compile(r'\b([CsgurtxyβαξλJ𝓘waFpK])_([a-z0-9]+)')
    for p in paragraphs:
        for run in list(p.runs):
            if not pattern.search(run.text):continue
            pieces=[];offset=0
            for m in pattern.finditer(run.text):pieces += [(run.text[offset:m.start()],False),(m[1],False),(m[2],True)];offset=m.end()
            pieces.append((run.text[offset:],False))
            for text,sub in pieces:
                if not text:continue
                e=deepcopy(run._r);r=Run(e,p);r.text=text
                if sub:r.font.subscript=True
                run._r.addprevious(e)
            run._r.getparent().remove(run._r)

if __name__=='__main__':
    t=(BASE/'q1_manuscript_v4.md').read_text(encoding='utf-8')
    roles='Conceptualization, Methodology, Software, Formal analysis, Investigation, Writing – original draft, Writing – review & editing, Visualization, Supervision.'
    t=t.replace('Author contributions. Romulo Sandoval Flórez, Claudia Patricia Parra Medina and German Leandro Vargas Fonseca reviewed the manuscript and take responsibility for its content. Individual CRediT role allocation remains to be supplied.', 'CRediT authorship contribution statement. '+ ' '.join(name+': '+roles for name in ['Romulo Sandoval Flórez','Claudia Patricia Parra Medina','German Leandro Vargas Fonseca'])+' All authors reviewed the manuscript and take responsibility for its content.')
    t=t.replace('Data availability. Source code, experiment settings, numerical records and figure-generation scripts are prepared for deposit. Public repository URL and persistent identifier: to be supplied.', 'Data availability. Source code, experiment settings, numerical records and analysis scripts are available at https://github.com/romusan/gears-inversion-reproducibility/tree/manuscript-v5. A SHA-256 manifest identifies the deposited files. No DOI has been assigned.')
    t=t.replace('- Noise comparisons use 100 realizations and a fixed observation budget',
        '- Minimum loaded flank coordinate screens relief visibility before inversion')
    t=t.replace('no two-load advantage is demonstrated by the deeper-support experiment.',
        'complementary-phase noise experiments show no statistically resolved two-load advantage in either support regime.')
    anchor='The present scope is quasistatic, ideal involute spur gears'
    discussion='The two approximate models give similar eight-phase ranges, 12.278 µm for FE and 11.662 µm for the baseline, while both exceed the published 10.527 µm amplitude. The deposited reference [1] defines STE from base radii and gear rotations, consistent with a line-of-action displacement, and uses a 0.21° angular increment in its three-dimensional FE calculation. Further comparison should audit angular sampling, tip and edge geometry, and the unmodified configuration. Agreement between approximate models does not identify the source of disagreement with the reference: they share geometric idealizations, the FE edge treatment is restricted, and its original single-pair sequence was not demonstrably asymptotic. The separated quadratic/distributed-pressure checks address interior numerical resolution; a converged continuous corner-contact cycle and the original reference settings remain necessary to resolve the physical discrepancy.\n\n'
    t=t.replace(anchor,discussion+anchor)
    t=t.replace('At s_d=0.40, local information and nonlinear recovery agree at low noise, and the tested 50/300 Nm allocation shows no demonstrated recovery advantage.',
        'Local information and nonlinear recovery agree at low noise in the deeper regime. A complementary-phase nonlinear control gives no statistically resolved advantage of 50/300 Nm at either relief start in the six tested noise cells.')
    t=t.replace('but a nonlinear recovery comparison on those controlled designs remains necessary.',
        'and the added complementary-phase reconstructions provide the nonlinear comparison for the tested phase assignment.')
    # Renumber main-table references from the unchanged source markers.
    t=v5.relabel(t)
    (BASE/'q1_manuscript_v5.md').write_text(t,encoding='utf-8')
    v5.SUPPLEMENT.clear()
    path=build(BASE/'q1_manuscript_v5.md',BASE/'MMT_manuscript_Q1_revision_v5.docx',BASE.parent/'referencias_40_q1.json',section_builder='review5_sections')
    d=Document(path)
    format_inline(d)
    for t in d.tables:
        for cell in t.rows[-2].cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    d.save(path)
    supplementary=Document(path)
    for child in list(supplementary.element.body):
        if child.tag!=qn('w:sectPr'):supplementary.element.body.remove(child)
    paragraph(supplementary,'Supplementary numerical controls for spur gear microgeometry inversion','Title')
    paragraph(supplementary,'Fisher diagnostics, load spectra, phase-grid controls and finite-element convergence for the accompanying manuscript. Tables distinguish the original statistical and linear nodal-contact records from the new complementary-phase and quadratic-element controls.')
    for r in sorted(v5.SUPPLEMENT,key=lambda r:int(r['number'][1:])):
        p=paragraph(supplementary,r['caption']);p.paragraph_format.keep_with_next=True
        table(supplementary,r['headers'],r['rows'],r['widths'])
    paragraph(supplementary,'S1 uses the original 1200-fit campaign. S2–S4 contain local information, not new Monte Carlo reconstructions. S5–S6 retain the original linear nodal FE results. S7 uses separate quadratic-displacement and distributed-pressure controls. The principal manuscript reports the 600 added nonlinear reconstructions and the limits of the corner-contact approximation.')
    format_inline(supplementary)
    for t in supplementary.tables:
        for cell in t.rows[-2].cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    supplementary.save(BASE/'MMT_supplementary_Q1_revision_v5.docx')
    print('Saved supplementary tables',len(v5.SUPPLEMENT))
