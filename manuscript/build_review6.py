"""Build revision 6 from the frozen v5 source and numerical section wrappers."""
from pathlib import Path
from copy import deepcopy
import json
from docx import Document
from docx.oxml.ns import qn
from build_corrected_manuscript import build,paragraph,table
from build_review5 import format_inline
import review5_sections as v5
BASE=Path(__file__).resolve().parent

def main():
    fe=json.loads((BASE.parent/'fe2d/results/quadratic_eight_phases.json').read_text())
    t=(BASE/'q1_manuscript_v5.md').read_text(encoding='utf-8')
    first='The two approximate models give similar eight-phase ranges, 12.278 µm for FE and 11.662 µm for the baseline, while both exceed the published 10.527 µm amplitude.'
    t=t.replace(first,f"The matched eight-phase range is {fe['sampled_range_um']:.3f} µm with quadratic FE and 11.662 µm with the baseline; both exceed the published 10.527 µm amplitude. The original linear FE range (12.278 µm) is retained as a discretization comparison, not the current FE estimate.")
    t=t.replace('Data availability. Source code, experiment settings, numerical records and analysis scripts are available at https://github.com/romusan/gears-inversion-reproducibility/tree/manuscript-v5. A SHA-256 manifest identifies the deposited files. No DOI has been assigned.',
        'Data availability. Source code, experiment settings, numerical records and analysis scripts are publicly available at https://github.com/romusan/gears-inversion-reproducibility/tree/manuscript-v6. The annotated tag manuscript-v6 and a SHA-256 manifest identify the deposited version. The preceding manuscript-v5 tag remains available. No DOI has been assigned.')
    (BASE/'q1_manuscript_v6.md').write_text(t,encoding='utf-8')
    v5.SUPPLEMENT.clear()
    path=build(BASE/'q1_manuscript_v6.md',BASE/'MMT_manuscript_Q1_revision_v6.docx',BASE.parent/'referencias_40_q1.json',section_builder='review6_sections')
    d=Document(path);format_inline(d)
    for tab in d.tables:
        for cell in tab.rows[-2].cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    d.save(path)
    s=Document(path)
    for child in list(s.element.body):
        if child.tag!=qn('w:sectPr'):s.element.body.remove(child)
    paragraph(s,'Supplementary numerical controls for spur gear microgeometry inversion','Title')
    paragraph(s,'Local information, phase-grid controls, finite-element resolution and angular sampling for the accompanying manuscript. The original linear FE records and the quadratic weak-contact records use different interface discretizations and are labelled separately.')
    for item in sorted(v5.SUPPLEMENT,key=lambda x:int(x['number'][1:])):
        p=paragraph(s,item['caption']);p.paragraph_format.keep_with_next=True
        table(s,item['headers'],item['rows'],item['widths'])
    paragraph(s,'S1–S4 retain the original statistical and local-information controls. S5–S6 and S8 preserve the original linear/nodal FE records. S7 varies quadratic bulk and pressure controls separately, with nonmonotonic interface stiffness. S9 reports all offsets for the declared interpolants; the eight-node FE waveform is an interpolation diagnostic and is not a direct FE angular scan.')
    format_inline(s)
    for tab in s.tables:
        for cell in tab.rows[-2].cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    s.save(BASE/'MMT_supplementary_Q1_revision_v6.docx')
    print('Saved revision 6, supplementary tables',len(v5.SUPPLEMENT))

if __name__=='__main__':main()
