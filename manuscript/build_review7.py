"""Build revision 7 from the frozen v6 markdown with controlled replacements.

v6 stays untouched; every edit is an asserted string replacement so an
unexpected v6 drift fails the build instead of silently patching the wrong text.
"""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from build_corrected_manuscript import build, paragraph, table
from build_review5 import format_inline
import review5_sections as v5
import review7_sections  # noqa: F401  (registers nothing; used by name)
BASE = Path(__file__).resolve().parent


def rep(t, old, new):
    assert old in t, 'v6 text drifted: ' + old[:80]
    return t.replace(old, new, 1)


def main():
    t = (BASE / 'q1_manuscript_v6.md').read_text(encoding='utf-8')

    # 1. Abstract: cross-operator confirmation sentence.
    t = rep(t,
        "Contact support thus explains internal observability, while independent stiffness and "
        "waveform evidence remains necessary for physical reconstruction.",
        "Contact support thus explains internal observability, while independent stiffness and "
        "waveform evidence remains necessary for physical reconstruction. An independent "
        "finite-element cycle confirms exact dedendum invisibility outside loaded support and "
        "yields addendum-depth recovery within 5% of its scale when the analytical operator "
        "inverts the finite-element response; repeating the nominal noise cells under the "
        "energy-projected compliance changes the median dedendum errors by less than 6%.")

    # 2. Highlights: one added bullet.
    t = rep(t,
        "- Mechanical ablations expose persistent external response discrepancies",
        "- Mechanical ablations expose persistent external response discrepancies\n"
        "- Independent finite-element cycles confirm exact relief invisibility outside loaded "
        "contact support")

    # 3. Section 2.2: adopt the energy-projected formulation for new experiments.
    t = rep(t,
        "No compliance factor is fitted to the external targets.",
        "No compliance factor is fitted to the external targets. For the robustness experiments "
        "of this revision the radial energy-projected variant is adopted as the reference "
        "formulation: it uses the shear modulus in the shear term, projects the contact force "
        "onto the section, and adds axial compression and the offset bending arm. At the nominal "
        "geometry and 300 Nm it raises peak-to-peak transmission error by 14% without changing "
        "contact support, singular structure or noise-free recovery. The legacy law remains the "
        "reproducible historical baseline.")

    # 4. Section 2.3: observation model carries no multiplicative term; E04 bias
    #    experiment described as unmodelled nuisance (archived records).
    t = rep(t,
        "The present independent-noise experiment uses the stated uncentred noise model.",
        "The present independent-noise experiment uses the stated uncentred noise model. The "
        "stated observation equation contains no multiplicative term; a gain miscalibration "
        "(1+γ) of the measurement chain is treated as an unmodelled nuisance. The archived bias "
        "experiment E04 (γ = 1%, 3% and 5% on otherwise noise-free data) shows the gain error "
        "being absorbed almost entirely by the weakly identified dedendum block: at γ=5% the "
        "estimated C_d moves by up to 0.045 mm and s_d by up to 0.14, while the addendum depth "
        "and start move by at most 0.0024 mm and 0.012. Unmodelled gain bias therefore "
        "contaminates exactly the coordinates that marginal-support identifiability already "
        "leaves unresolved.")

    # 5. New subsection 4.5 before the Discussion.
    t = rep(t, "## 5 Discussion and limitations",
        "### 4.5 Cross-operator inversion and formulation robustness\n\n"
        "Q1_CROSS_INVERSE\n\n"
        "## 5 Discussion and limitations")

    # 6. Discussion: the cross-operator test partially fills the previously
    #    identified evidence gap and bounds practical precision.
    t = rep(t,
        "An independently generated contact/TE dataset with known geometry, load sharing and "
        "sufficient angular resolution is needed before physical parameter-recovery claims.",
        "An independently generated contact/TE dataset with known geometry, load sharing and "
        "sufficient angular resolution is needed before physical parameter-recovery claims. The "
        "cross-operator test of Section 4.5 partially fills that gap: finite-element responses "
        "of known geometry inverted with the analytical operator recover the addendum depth "
        "within 5% of its scale, reproduce exact dedendum invisibility outside loaded support, "
        "and bound the residual between the two mechanical operators at 0.109–0.151 µm of "
        "transmission error. The systematic external peak-to-peak discrepancy remains "
        "unresolved, and spurious small dedendum depths show that discrepancy absorption, not "
        "physical relief, can dominate that block in cross-model fits.")

    # 7. Conclusions: bounded cross-operator statement.
    t = rep(t,
        "Additional edge-bearing phases cannot establish full-cycle benchmark agreement.",
        "Additional edge-bearing phases cannot establish full-cycle benchmark agreement. The "
        "cross-operator test inverts finite-element responses with the analytical operator: the "
        "addendum depth is recovered within 5% of its scale, no dedendum signal exists to "
        "recover outside loaded support, and the inter-operator residual bounds practical "
        "precision at about 0.1 µm of transmission error. The benchmark amplitude discrepancy "
        "itself remains structural and unresolved.")

    # 8. Data availability: v7 tag.
    t = rep(t,
        "https://github.com/romusan/gears-inversion-reproducibility/tree/manuscript-v6. The "
        "annotated tag manuscript-v6 and a SHA-256 manifest identify the deposited version. The "
        "preceding manuscript-v5 tag remains available. No DOI has been assigned.",
        "https://github.com/romusan/gears-inversion-reproducibility/tree/manuscript-v7. The "
        "annotated tag manuscript-v7 and a SHA-256 manifest identify the deposited version. The "
        "preceding manuscript-v6 tag remains available. No DOI has been assigned.")

    (BASE / 'q1_manuscript_v7.md').write_text(t, encoding='utf-8')

    v5.SUPPLEMENT.clear()
    path = build(BASE / 'q1_manuscript_v7.md',
                 BASE / 'MMT_manuscript_Q1_revision_v7.docx',
                 BASE.parent / 'referencias_40_q1.json',
                 section_builder='review7_sections', math_revision=7)
    d = Document(path); format_inline(d)
    for tab_ in d.tables:
        for cell in tab_.rows[-2].cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    d.save(path)

    s = Document(path)
    for child in list(s.element.body):
        if child.tag != qn('w:sectPr'):
            s.element.body.remove(child)
    paragraph(s, 'Supplementary numerical controls for spur gear microgeometry inversion', 'Title')
    paragraph(s, 'Local information, phase-grid controls, finite-element resolution and angular '
                 'sampling for the accompanying manuscript. The original linear FE records and '
                 'the quadratic weak-contact records use different interface discretizations and '
                 'are labelled separately. Tables S10–S11 report the cross-operator inversion '
                 'and formulation-robustness records of revision 7.')
    for item in sorted(v5.SUPPLEMENT, key=lambda x: int(x['number'][1:])):
        p = paragraph(s, item['caption']); p.paragraph_format.keep_with_next = True
        table(s, item['headers'], item['rows'], item['widths'])
    paragraph(s, 'S1–S4 retain the original statistical and local-information controls. S5–S6 and '
                 'S8 preserve the original linear/nodal FE records. S7 varies quadratic bulk and '
                 'pressure controls separately, with nonmonotonic interface stiffness. S9 reports '
                 'all offsets for the declared interpolants; the eight-node FE waveform is an '
                 'interpolation diagnostic and is not a direct FE angular scan. S10 reports '
                 'cross-operator inversion against the two-dimensional finite-element generator '
                 '(nineteen interior phases; six restricted phases excluded). S11 repeats the '
                 'nominal noise cells under the energy-projected compliance formulation.')
    format_inline(s)
    for tab_ in s.tables:
        for cell in tab_.rows[-2].cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    s.save(BASE / 'MMT_supplementary_Q1_revision_v7.docx')
    print('Saved revision 7, supplementary tables', len(v5.SUPPLEMENT))


if __name__ == '__main__':
    main()
