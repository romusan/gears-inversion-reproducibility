import re,json
from pathlib import Path
from docx import Document
from docx.shared import Pt,RGBColor
from build_corrected_manuscript import paragraph
BASE=Path(__file__).resolve().parent;ROOT=BASE.parent
def final_bibliography(doc):
    refs=json.loads((ROOT/'referencias_42_v12.json').read_text(encoding='utf-8'))
    for p in doc.paragraphs:
        match=re.match(r'^\[(\d+)\]',p.text)
        if not match:continue
        n=int(match[1]);r=refs[n-1]
        def name(a):
            initials=''.join(x[0]+'.' for x in re.findall(r'[^\W\d_]+',a.get('given',''),flags=re.UNICODE))
            return (initials+' '+a.get('family','')).strip()
        authors=', '.join(name(a) for a in r.get('author',[])[:6])
        if len(r.get('author',[]))>6:authors+=', et al.'
        title=r['title'][0].rstrip('.');journal=(r.get('container-title') or [''])[0]
        year=r.get('year') or r['published']['date-parts'][0][0]
        if r.get('type') in ['book','monograph']:
            body=f"{authors}, {title}, {r.get('publisher','')}, {year}."
        elif r.get('type')=='standard':
            body=f"{authors}, {title}, ISO, Geneva, {year}."
        else:
            volume=r.get('volume') or '';pages=r.get('page') or r.get('article-number') or ''
            body=f"{authors}, {title}, {journal} {volume} ({year}) {pages}."
        link='https://doi.org/'+r['DOI'] if r.get('DOI') else r.get('url','')
        p.text=f'[{n}] {body} {link}'
        for run in p.runs:run.font.size=Pt(10)
def write_highlights():
    doc=Document();paragraph(doc,'Highlights','Title')
    doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Title'].font.name='Times New Roman'
    doc.styles['Title'].font.size=Pt(16)
    for element in [doc.element,doc.styles.element]:
        for border in list(element.xpath('.//w:pBdr')):border.getparent().remove(border)
    for run in doc.paragraphs[0].runs:run.font.color.rgb=RGBColor(0,0,0)
    lines=(BASE/'highlights_v12.txt').read_text(encoding='utf-8').splitlines()
    bullets=[x[2:] for x in lines if x.startswith('• ')]
    assert len(bullets)==5 and all(len(x)<=85 for x in bullets)
    for text in bullets:paragraph(doc,'• '+text)
    doc.save(BASE/'Highlights_v12.docx')
