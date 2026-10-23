"""Insert verified numerical results into the next article revision."""
import json
from pathlib import Path
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'gear_inverse_identifiability'/'results'/'q1_revision_2026-10-12'

def read(name):return json.loads((OUT/name).read_text(encoding='utf-8'))
def pic(doc,name,caption,paragraph):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(OUT/name),width=Cm(15.5))
    p=paragraph(doc,caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(9)

def add_section(doc,marker,paragraph,table,figure):
    if marker=='Q1_BENCHMARK':
        report=read('forward_ablation.json');rr=report['results'];geometry=report['geometry']
        paragraph(doc,f"The computed working pressure angle is {geometry['alpha_w_deg']:.3f}° and transverse contact ratio is {geometry['contact_ratio']:.4f}, compared with the reference's rounded 1.45. Table 1 reports signed absolute differences; Figure 1 also retains case-relative errors. Normalizing baseline differences by the unmodified FE peak-to-peak value gives {min(r['unmodified_normalized_percent'] for r in rr if r['variant']=='legacy'):.1f}% to {max(r['unmodified_normalized_percent'] for r in rr if r['variant']=='legacy'):.1f}% across modified and unmodified cases.")
        data=[]
        for name in ['unmodified','Sol1','Sol2','Sol3']:
            for T in [350,500,650]:
                rows={r['variant']:r for r in rr if r['case']==name and r['torque_Nm']==T}
                data.append([name,T,f"{rows['legacy']['reference_fe_pp_um']:.3f}",*[f"{rows[v]['absolute_difference_um']:+.3f}" for v in ['legacy','shear-G','projected-energy','centerline-energy']]])
        paragraph(doc,'Table 1 FE peak-to-peak targets and signed model-minus-FE differences in µm at 240 angular positions. All twelve cases exceed the internal 5% target for each variant; surrogate predictions are excluded.')
        table(doc,['Case','Nm','FE µm','Baseline Δ','G shear Δ','Radial energy Δ','Centre-line Δ'],data,[2.6,1.3,1.8,2.4,2.4,2.7,3.2])
        for variant,label in [('legacy','Baseline'),('shear-G','G-only'),('projected-energy','Radial energy'),('centerline-energy','Centre-line energy')]:
            r=[t for t in rr if t['variant']==variant]
            paragraph(doc,f"{label}: case-relative differences span {min(t['difference_percent'] for t in r):.1f}–{max(t['difference_percent'] for t in r):.1f}%. Signed peak-to-peak errors span {min(t['absolute_difference_um'] for t in r):.3f}–{max(t['absolute_difference_um'] for t in r):.3f} µm. These are systematic peak-to-peak discrepancies, not measured waveform-noise standard deviations.")
        pic(doc,'forward_ablation.png','Figure 1 Physical compliance ablations compared with the same external FE targets. The shaded ±5% band is the internal project target.',paragraph)
    elif marker=='Q1_ANGULAR':
        report=read('angular_peak_convergence.json');rows=report['rows']
        for name in ['unmodified','Sol2']:
            rr=[r for r in rows if r['variant']=='legacy' and r['case']==name]
            paragraph(doc,f"At 650 Nm, baseline {name} peak-to-peak TE changes from {next(r['peak_to_peak_um'] for r in rr if r['n_positions']==120):.4f} µm at 120 positions to {rr[-1]['peak_to_peak_um']:.4f} µm at 1920 positions; the last 960-to-1920 change is {rr[-1]['relative_pp_change']*100:.3f}%. Refinement is not monotonic because the peak lies near a contact transition. A small cycle-RMS change therefore does not alone certify peak or derivative convergence.")
        pic(doc,'angular_peak_convergence.png','Figure 2 Angular peak refinement and detail of the baseline unmodified TE curve at 650 Nm. Closely spaced points resolve narrow transitions that coarser sampling misses.',paragraph)
    elif marker=='Q1_TRANSITION':
        rows=read('support_transition.json')['rows']
        invisible=[r for r in rows if r['dedendum_signal_rms_um']<1e-10]
        selected=[min(rows,key=lambda r:abs(r['sd_true']-v)) for v in [.25,.4,.5]]
        paragraph(doc,f"The sampled starts {invisible[0]['sd_true']:.2f}–{invisible[-1]['sd_true']:.2f} give zero dedendum signal and numerical rank two. At starts 0.25,0.40 and 0.50 the signal RMS is {selected[0]['dedendum_signal_rms_um']:.6f}, {selected[1]['dedendum_signal_rms_um']:.6f} and {selected[2]['dedendum_signal_rms_um']:.6f} µm, respectively, with rank four. It first exceeds 0.1 µm at the tested start 0.43. This per-sample comparison is descriptive; accumulated information depends on the sample count and parameter coupling.")
        paragraph(doc,'The near-invisible nominal setting is therefore not representative of every admissible dedendum. Its behaviour follows marginal penetration into loaded contact, and the recoverable deeper-relief regime is retained as a counterexample to a universal ambiguity claim.')
        pic(doc,'support_transition.png','Figure 3 Transition in dedendum signal, local rank and noiseless multistart recovery as the true relief start moves into the loaded flank.',paragraph)
    elif marker=='Q1_PROFILES':
        rr=read('profile_likelihood.json')['records']
        assert len(rr)==8,'Incomplete likelihood profiles'
        for sd in [.25,.4]:
            for parameter in ['sd','Cd']:
                summaries=[]
                for rec in rr:
                    if rec['sd_true']==sd and rec['parameter']==parameter:
                        accepted=[r['value'] for r in rec['rows'] if r['delta_chi2']<=3.841]
                        factor=1000 if parameter=='Cd' else 1
                        label='300 Nm' if rec['design']=='single300' else '50/300 Nm'
                        summaries.append(f"{label}: [{min(accepted)*factor:.3g},{max(accepted)*factor:.3g}]" if accepted else f"{label}: no grid point below reference")
                paragraph(doc,f"For true s_d={sd:.2f}, the sampled {parameter} profile values below the nominal Δχ² reference are "+'; '.join(summaries)+(' µm.' if parameter=='Cd' else '.')+' These grid sets are diagnostics, not coverage-certified confidence intervals.')
        failures=sum(not r['success'] for rec in rr for r in rec['rows'])
        paragraph(doc,f"Profile-point optimizations report {failures} selected nonconverged solutions. Their recorded objective values are retained; a stalled point can only overstate the profiled cost and requires caution when interpreting a narrow set.")
        pic(doc,'profile_likelihood.png','Figure 4 Dedendum profile objectives with the other three coordinates refitted. Curves are clipped vertically at Δχ²=15 for readability; the complete objectives and convergence flags accompany the data.',paragraph)
    elif marker=='Q1_MONTE_CARLO':
        report=read('monte_carlo_summary.json');rows=report['summary']
        assert report['records']==1200 and all(r['n']==100 for r in rows),'Incomplete Monte Carlo campaign'
        success=sum(r['success'] for r in rows)
        paragraph(doc,f"The campaign contains {report['records']} selected reconstructions across twelve cells, with 100 seeds per cell and {success} selected fits reporting optimizer convergence. A successful termination is distinct from attaining a global optimum; both-start termination and objective values are available in the records. Table 2 includes all selected fits, including nonconverged ones.")
        paragraph(doc,'Table 2 Monte Carlo medians at fixed total count. Parameter errors use the declared scales; profile and held-out prediction errors are in µm. Success is the number of selected fits reporting convergence out of 100.')
        data=[]
        for r in rows:
            data.append([r['sd_true'],'300' if r['design']=='single300' else '50/300',r['sigma_um'],r['success'],
                         f"{r['addendum_scaled_error']['median']:.3g}",f"{r['dedendum_scaled_error']['median']:.3g}",
                         f"{r['full_profile_rmse_um']['median']:.3g}",f"{r['heldout200_rmse_um']['median']:.3g}"])
        table(doc,['True s_d','Nm','σ µm','Success','Addendum','Dedendum','Profile RMS','Held-out RMS'],data,[1.6,1.6,1.5,1.6,2.1,2.1,2.95,2.95])
        pic(doc,'monte_carlo.png','Figure 5 Dedendum error under 100 noise realizations per cell. Solid lines show medians; dotted lines show empirical 90th percentiles. Both designs have 120 observations.',paragraph)
        paired=read('paired_comparison.json')['records']
        for r in paired:
            m=r['dedendum_scaled_error'];lo,hi=m['bootstrap_percentile_95']
            paragraph(doc,f"For true s_d={r['sd_true']:.2f}, σ={r['sigma_um']:.1f} µm, the median paired dedendum-error difference (two loads minus single load) is {m['median_paired_difference']:.3g}, with percentile-bootstrap bounds [{lo:.3g},{hi:.3g}] from 5000 resamples of the 100 seed pairs. These bounds characterize uncertainty in the simulation statistic, not physical parameter intervals.")
    elif marker=='Q1_SPECTRA':
        spectra=read('support_transition.json')['spectra']
        for sd in [.25,.4]:
            rr=[r for r in spectra if r['sd_true']==sd]
            for j in [2,3]:
                vals=[r['singular_values'][j] for r in rr]
                paragraph(doc,f"At true s_d={sd}, singular value {j+1} ranges from {min(vals):.4g} to {max(vals):.4g} over 50–300 Nm. This checks the near-load-invariance hypothesis directly rather than assuming visually coincident curves are exactly identical.")
        pic(doc,'analytic_spectra.png','Figure 6 Full singular spectra from the implicit tangent at two relief starts. Integer indices and torque labels identify the four directions without implying that the singular vectors are fixed across loads.',paragraph)
    elif marker=='Q1_GEOMETRIES':
        rr=read('geometry_transfer.json')['records']
        paragraph(doc,'Additional ideal involute pairs are G1=(20,60,2.5,15), G2=(30,45,4.0,25) and G3=(18,64,2.0,12), with fields (z₁,z₂,m mm,b mm); all use α=20°. The 18-tooth pinion replaces the earlier 16-tooth case, avoiding the standard unshifted undercut threshold 2/sin²20°≈17.1. Relief depths scale with m, and every geometry uses the same physical bounds C_a≤0.100 mm and C_d≤0.050 mm. This is an ideal involute portability test with an approximate root, not manufacturing validation.')
        paragraph(doc,'Table 3 Noise-free geometry transfer under two true starts and equal observation counts. Error is max absolute parameter error divided by the declared scales; fit residual is in µm.')
        data=[];scale=[.05,1,.01,1]
        for r in rr:
            errors=[abs(a-b)/s for a,b,s in zip(r['theta'],r['theta_true'],scale)]
            data.append([r['geometry'],r['sd_true'],'/'.join(str(int(T)) for T in r['torques_Nm']),f"{max(errors[:2]):.3g}",f"{max(errors[2:]):.3g}",f"{r['fit_rmse_um']:.3g}"])
        table(doc,['Geometry','True s_d','Nm','Addendum','Dedendum','Fit µm'],data,[2.3,2.1,2.4,3.2,3.2,3.2])
    else:raise ValueError(marker)
