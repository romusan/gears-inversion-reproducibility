"""Second-review tables; retain the first revision's section builder."""
import json
from pathlib import Path
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from q1_sections import add_section as previous_section
ROOT=Path(__file__).resolve().parent.parent
OLD=ROOT/'gear_inverse_identifiability/results/q1_revision_2026-10-12'
NEW=OLD.parent/'q1_review2_2026-10-12'
LABELS={'legacy':'Baseline','shear-G':'G shear','projected-energy':'Radial energy','centerline-energy':'Centre-line'}
def read(name,new=False):return json.loads(((NEW if new else OLD)/name).read_text(encoding='utf-8'))
def loads(name):return '300' if name=='single300' else '50/300'
def pic(doc,name,caption,paragraph,new=False):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str((NEW if new else OLD)/name),width=Cm(15.5))
    p=paragraph(doc,caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(9)
def tab(doc,num,caption,headers,rows,widths,paragraph,table):
    paragraph(doc,f'Table {num} '+caption);table(doc,headers,rows,widths)

def add_section(doc,marker,paragraph,table,figure):
    if marker=='Q1_BENCHMARK':
        report=read('forward_ablation.json');rr=report['results'];g=report['geometry']
        paragraph(doc,f"Working pressure angle is {g['alpha_w_deg']:.3f}° and contact ratio {g['contact_ratio']:.4f}, compared with the rounded reference ratio 1.45. All 48 model-minus-FE peak-to-peak differences are positive. Baseline differences normalized by the unmodified FE amplitude span {min(r['unmodified_normalized_percent'] for r in rr if r['variant']=='legacy'):.1f}–{max(r['unmodified_normalized_percent'] for r in rr if r['variant']=='legacy'):.1f}%. The systematic sign indicates response discrepancy, but does not specify waveform-noise covariance.")
        rows=[]
        for name in ['unmodified','Sol1','Sol2','Sol3']:
            for T in [350,500,650]:
                r={x['variant']:x for x in rr if x['case']==name and x['torque_Nm']==T}
                rows.append([name,T,f"{r['legacy']['reference_fe_pp_um']:.3f}",*[f"{r[v]['absolute_difference_um']:+.3f}" for v in LABELS]])
        tab(doc,1,'FE peak-to-peak targets and signed differences in µm at 240 positions. Each variant fails the adopted 5% target in all twelve cases.',
            ['Case','Nm','FE µm','Baseline Δ','G shear Δ','Radial energy Δ','Centre-line Δ'],rows,[2.6,1.3,1.8,2.4,2.4,2.7,3.2],paragraph,table)
        pic(doc,'forward_ablation.png','Figure 1 Compliance variants compared with external FE amplitudes. The shaded ±5% band is the adopted research target.',paragraph,True)
    elif marker=='Q1_ANGULAR':
        old=read('angular_peak_convergence.json')['rows'];wave=read('waveform_diagnostics.json',True);forward=read('forward_ablation.json')['results']
        paragraph(doc,'Correcting the shear denominator increases the benchmark discrepancy. Table 2 contrasts error ranges with internal approach levels and incremental mesh stiffness. Approach levels are conditional medians for one or two loaded pairs. Stiffness is d(sum p_k)/dδ at fixed gaps, rather than force/approach. External FE waveform, plateau and stiffness data are unavailable for these cases; these diagnostics therefore provide no additional fidelity validation.')
        rows=[]
        for v,label in LABELS.items():
            f=[r for r in forward if r['variant']==v];w=next(r for r in wave['rows'] if r['variant']==v and r['case']=='unmodified')
            rows.append([label,f"{min(r['difference_percent'] for r in f):.1f}–{max(r['difference_percent'] for r in f):.1f}",f"{min(r['absolute_difference_um'] for r in f):.3f}–{max(r['absolute_difference_um'] for r in f):.3f}",f"{w['median_single_pair_approach_um']:.2f}",f"{w['median_two_pair_approach_um']:.2f}",f"{w['waveform_rms_um']:.2f}",f"{w['mesh_stiffness_min_N_um']:.1f}–{w['mesh_stiffness_max_N_um']:.1f}"])
        tab(doc,2,'Error ranges over twelve benchmark cases and internal unmodified response at 650 Nm, 1920 positions. Approach and RMS are µm; stiffness is N/µm.',
            ['Variant','Relative %','Signed Δ µm','1 pair','2 pairs','RMS','Stiffness'],rows,[2.6,2.2,2.2,1.8,1.8,1.5,4.3],paragraph,table)
        checks=[r for r in old if r['variant']=='legacy' and r['n_positions']==1920]
        paragraph(doc,f"At 650 Nm, baseline peak-to-peak changes from 960 to 1920 positions are {checks[0]['relative_pp_change']*100:.3f}% without relief and {checks[1]['relative_pp_change']*100:.3f}% for Sol2. Nonmonotonic refinement near narrow contact transitions prevents cycle-RMS convergence from certifying every peak or derivative.")
        local=wave['transition_refinement']
        paragraph(doc,'A local 0.06° window resolves the steep baseline transition: '+', '.join(f"Δφ={r['step_deg']:.3g}° gives max adjacent change {r['max_adjacent_change_um']:.4g} µm" for r in local)+'. Dividing the step by four reduces the largest change by about four, consistent with a narrow continuous transition rather than a discontinuity. Exact corner geometry and body deformation could change its shape; that physical difference is not quantified here.')
        pic(doc,'plateau_stiffness.png','Figure 2 Internal total approach and incremental mesh stiffness for the unmodified benchmark at 650 Nm. Variant names match the tables.',paragraph,True)
    elif marker=='Q1_TRANSITION':
        previous_section(doc,marker,paragraph,table,figure)
    elif marker=='Q1_PROFILES':
        profiles=read('profile_likelihood.json')['records'];assert len(profiles)==8;rows=[]
        for r in profiles:
            factor=1000 if r['parameter']=='Cd' else 1;accepted=[p['value']*factor for p in r['rows'] if p['delta_chi2']<=3.841]
            rows.append([r['sd_true'],loads(r['design']),'C_d µm' if factor==1000 else 's_d',f"{(r['rows'][1]['value']-r['rows'][0]['value'])*factor:.3g}",f"{min(accepted):.3g}–{max(accepted):.3g}",f"{max(p['delta_chi2'] for p in r['rows']):.3g}"])
        tab(doc,3,'Profile grid resolution and accepted-grid extrema at Δχ²≤3.841. Values are sampled extrema, not interpolated or coverage-certified confidence endpoints.',
            ['True s_d','Nm','Parameter','Step','Grid extrema','Max Δχ²'],rows,[1.6,1.8,2.6,1.8,4.2,4.4],paragraph,table)
        cert=read('flat_profile_diagnostics.json',True)['rows'];m=[r for r in cert if r['sd_true']==.25];d=[r for r in cert if r['sd_true']==.4]
        paragraph(doc,f"At s_d=0.25, profiles are nearly flat across both full domains. Fitting the exactly invisible subset C_d=0, s_d=0.05 increases χ² by {m[0]['invisible_delta_chi2']:.3g} and {m[1]['invisible_delta_chi2']:.3g} for 300 and 50/300 Nm. At this start, changing C_d from zero to 50 µm leaves TE unchanged under the support condition. The noisy data therefore cannot resolve departure from structurally invisible sets; this is distinct from declaring the nonzero-sensitivity truth structurally unidentifiable. At s_d=0.40, that subset incurs Δχ²={d[0]['invisible_delta_chi2']:.2f} and {d[1]['invisible_delta_chi2']:.2f}, above the nominal reference. All selected profile-point fits converge.")
        pic(doc,'profile_likelihood.png','Figure 4 Profiles with other coordinates refitted. Values above Δχ²=15 are clipped only for display; grid steps are 0.01 in start and 1.25 µm in depth.',paragraph)
    elif marker=='Q1_MONTE_CARLO':
        report=read('monte_carlo_summary.json');summary=report['summary'];diag=read('fit_diagnostics.json',True)
        assert report['records']==1200 and all(r['n']==100 for r in summary)
        paragraph(doc,f"The twelve cells contain 1200 selected reconstructions, with {sum(r['success'] for r in summary)} reporting convergence. All selected fits, including the 45 failures, enter the summaries; termination does not establish a global minimum. Tables 4 and 5 separate recovery from bound and starting-point diagnostics.")
        rows=[[r['sd_true'],loads(r['design']),r['sigma_um'],r['success'],f"{r['addendum_scaled_error']['median']:.3g}",f"{r['dedendum_scaled_error']['median']:.3g}",f"{r['full_profile_rmse_um']['median']:.3g}",f"{r['heldout200_rmse_um']['median']:.3g}"] for r in summary]
        tab(doc,4,'Monte Carlo medians at 120 observations. Parameter errors use declared scales; profile and held-out errors are µm. Success is out of 100 seeds.',
            ['True s_d','Nm','σ µm','Success','Addendum','Dedendum','Profile RMS','Held-out RMS'],rows,[1.6,1.6,1.5,1.6,2.1,2.1,2.95,2.95],paragraph,table)
        paragraph(doc,'The two starts have baseline scaled addendum errors 0.20 and 0.10, and dedendum errors 0.50 and 1.50, at either truth. A flat objective permits dependence on initialization and box constraints, but nonmonotonic error medians do not prove that estimates stayed at their starts. Actual movement from the selected start is reported below. Maximum dedendum error is capped at 4.5 when C_d reaches 50 µm; 90th-percentile saturation in Figure 5 therefore has a box-constraint interpretation.')
        rows=[[r['sd_true'],loads(r['design']),r['sigma_um'],f"{100*r['dedendum_bound_fraction']:.0f}",f"{100*r['cd_upper_fraction']:.0f}",f"{100*r['sd_any_bound_fraction']:.0f}",f"{r['median_distance_from_selected_start_dedendum']:.3g}"] for r in diag['fit_bound_statistics']]
        tab(doc,5,'Bound-hit percentages and median maximum scaled dedendum movement from the selected start. A hit is within 10⁻⁶ of box width from a bound; all selected fits are included.',
            ['True s_d','Nm','σ µm','Any ded. %','C_d upper %','s_d bound %','Movement'],rows,[1.6,1.7,1.6,2.6,2.8,2.8,3.3],paragraph,table)
        pic(doc,'monte_carlo.png','Figure 5 Dedendum error under 100 realizations per cell. Solid lines are medians; dotted lines are 90th percentiles and can reach the imposed error cap 4.5.',paragraph)
        rows=[]
        for r in read('paired_comparison.json')['records']:
            m=r['dedendum_scaled_error'];lo,hi=m['bootstrap_percentile_95']
            rows.append([r['sd_true'],r['sigma_um'],f"{m['median_paired_difference']:+.4g}",f"{lo:+.4g}",f"{hi:+.4g}",f"{100*m['fraction_multi_lower']:.0f}"])
        tab(doc,6,'Paired two-load-minus-single-load dedendum error, with 5000 bootstrap resamples of 100 seed pairs. Negative favours 50/300 Nm. Bounds concern a simulation statistic, not physical parameter intervals.',
            ['True s_d','σ µm','Median Δ','Lower 2.5%','Upper 97.5%','Pair better %'],rows,[1.8,1.8,3.0,3.0,3.0,3.8],paragraph,table)
        paragraph(doc,'All six intervals contain zero, and five of six paired medians favour a single 300 Nm load. There is no demonstrable advantage of the tested two-load allocation. This does not prove equivalence or exclude a pair that changes contact support more effectively.')
        rows=[[r['sd_true'],loads(r['design']),f"{r['singular_values'][-1]:.4g}",f"{r['gaussian_abs_median']:.4g}",f"{r['mc_projected_abs_median']:.4g}",f"{r['gaussian_dedendum_max_median']:.4g}",f"{r['mc_dedendum_max_median']:.4g}"] for r in diag['fisher_comparison']]
        tab(doc,7,'Fisher–Monte Carlo comparison at σ=0.1 µm. Predicted weakest-direction absolute median is 0.67449/σ₄; projected MC uses the same singular vector. The dedendum maximum is a distinct statistic approximated by 100000 draws from the full unconstrained Gaussian covariance.',
            ['True s_d','Nm','σ₄','Local dir.','MC dir.','Local ded. max','MC ded. max'],rows,[1.6,1.8,1.8,2.4,2.4,3.2,3.2],paragraph,table)
        paragraph(doc,'At s_d=0.40 and 300 Nm, weakest-direction standard deviation 1/σ₄≈0.157 gives absolute median 0.1059, versus 0.1075 in projected Monte Carlo. The Gaussian dedendum-maximum median 0.1041 also agrees with observed 0.1058. At s_d=0.25, unconstrained local Gaussian medians extend far beyond box-limited errors, so nonlinear profiles rather than Fisher covariance describe the unresolved coordinates.')
    elif marker=='Q1_SPECTRA':
        rows=[[r['sd_true'],int(r['torque_Nm']),f"{r['singular_values'][2]:.4g}",f"{r['singular_values'][3]:.4g}"] for r in read('support_transition.json')['spectra']]
        tab(doc,8,'Third and fourth scaled whitened singular values at 120 observations, σ=0.1 µm. Near-invariance with load occurs only at marginal support; singular vectors need not remain fixed.',
            ['True s_d','Torque Nm','σ₃','σ₄'],rows,[3.8,3.8,4.4,4.4],paragraph,table)
        pic(doc,'analytic_spectra.png','Figure 6 Full singular spectra at marginal and deeper support. Torque labels and integer indices match Table 8.',paragraph)
    elif marker=='Q1_GEOMETRIES':
        paragraph(doc,'The nominal geometry is G0=(24,48,3.0,20); additional pairs are G1=(20,60,2.5,15), G2=(30,45,4.0,25) and G3=(18,64,2.0,12), with fields (z₁,z₂,m mm,b mm). All have α=20° and pinion count above the unshifted undercut threshold 2/sin²20°≈17.1. Truth depths scale with module; bounds remain C_a≤0.100 mm and C_d≤0.050 mm. Table 9 connects geometry with loaded support and information rather than treating a machine-precision self-fit as mechanical validation.')
        rows=[[r['geometry'],r['sd_true'],'/'.join(str(int(T)) for T in r['torques_Nm']),f"{r['min_loaded_xi']:.4f}",r['loaded_relief_contacts'],f"{r['smallest_singular_value']:.3g}",r['rank'],f"{r['dedendum_scaled_error']:.3g}"] for r in read('geometry_support.json',True)['rows']]
        tab(doc,9,'Support and sensitivity by geometry and design: 120 observations, σ=0.1 µm for scaling. Count includes repeated phases and pairs; dedendum error is scaled noiseless recovery error.',
            ['Pair','True s_d','Nm','ξ_min','ξ<s_d count','σ₄','Rank','Ded. error'],rows,[1.6,1.6,1.8,2.0,2.6,2.3,1.3,3.2],paragraph,table)
        paragraph(doc,'For G2 at s_d=0.25, minimum loaded ξ exceeds the relief start, dedendum Jacobian columns vanish and rank is two. A large dedendum error therefore coexists with a self-consistent TE fit. At s_d=0.40, loaded support enters the relief and four-parameter recovery is possible in the noise-free test. This explains the G2 contrast without attributing it to differing angular sampling.')
    else:raise ValueError(marker)
